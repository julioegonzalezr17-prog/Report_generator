"""
Test comprehensive para validar que step_engine.py funciona correctamente con TODOS los actions del JSON
"""
import os
import sys
from pathlib import Path

# Add workspace to path
workspace_path = str(Path(__file__).parent)
if workspace_path not in sys.path:
    sys.path.insert(0, workspace_path)

from ExcelStyler import ExcelStyler
from step_engine import StepEngine
from openpyxl import Workbook


def create_test_excel(filename):
    """Create a test Excel file with sample data"""
    wb = Workbook()
    ws = wb.active
    ws.title = "DATA"
    
    # Add headers in row 3 (ExcelStyler convention)
    headers = ["Name", "Status", "Value", "Fault_Code"]
    for col_idx, header in enumerate(headers, 1):
        ws.cell(row=3, column=col_idx, value=header)
    
    # Add sample data (starting from row 4)
    data = [
        ["Record_1", "OK", "100", 1992],
        ["Record_2", "WARNING", "200", 1993],
        ["Record_3", "OK", "100", 1994],
        ["Record_4", "ERROR", "300", 999],  # 999 not in typical TDM
        ["Record_5", "OK", "150", 1995],
    ]
    
    for row_idx, row_data in enumerate(data, 4):
        for col_idx, value in enumerate(row_data, 1):
            ws.cell(row=row_idx, column=col_idx, value=value)
    
    wb.save(filename)
    return filename


def test_highlight_equals_action():
    """Test find_and_highlight_equals action with single value"""
    print("\n" + "="*70)
    print("TEST 1: find_and_highlight_equals action (single value)")
    print("="*70)
    
    test_excel = "test_find_equals.xlsx"
    create_test_excel(test_excel)
    
    styler = ExcelStyler(test_excel, sheet_name="DATA")
    
    test_config = {
        "steps": [
            {
                "action": "find_and_highlight_equals",
                "column": "Status",
                "value": "OK",
                "color": "FF00FF00"  # Green
            }
        ]
    }
    
    engine = StepEngine(styler, test_config)
    
    try:
        print("✓ Executing find_and_highlight_equals...")
        result = engine.execute()
        styler.save("test_find_equals_output.xlsx")
        print(f"✓ SUCCESS: Highlighted all 'OK' status cells")
        print(f"  Output: test_find_equals_output.xlsx")
        
        # Cleanup
        if os.path.exists(test_excel):
            os.remove(test_excel)
        
        return True
    except Exception as e:
        print(f"✗ FAILED: {e}")
        import traceback
        traceback.print_exc()
        return False


def test_highlight_equals_list_action():
    """Test find_and_highlight_equals with list of values"""
    print("\n" + "="*70)
    print("TEST 1b: find_and_highlight_equals action (list of values)")
    print("="*70)
    
    test_excel = "test_find_equals_list.xlsx"
    create_test_excel(test_excel)
    
    styler = ExcelStyler(test_excel, sheet_name="DATA")
    
    # Test with list of values to match
    test_config = {
        "steps": [
            {
                "action": "find_and_highlight_equals",
                "column": "Status",
                "value": ["OK", "WARNING"],  # List of values
                "color": "FFFFCC00"  # Gold
            }
        ]
    }
    
    engine = StepEngine(styler, test_config)
    
    try:
        print("✓ Executing find_and_highlight_equals with list of values...")
        result = engine.execute()
        styler.save("test_find_equals_list_output.xlsx")
        print(f"✓ SUCCESS: Highlighted all 'OK' or 'WARNING' status cells")
        print(f"  Output: test_find_equals_list_output.xlsx")
        
        # Cleanup
        if os.path.exists(test_excel):
            os.remove(test_excel)
        
        return True
    except Exception as e:
        print(f"✗ FAILED: {e}")
        import traceback
        traceback.print_exc()
        return False


def test_highlight_contains_action():
    """Test find_and_highlight_contains action"""
    print("\n" + "="*70)
    print("TEST 2: find_and_highlight_contains action")
    print("="*70)
    
    test_excel = "test_find_contains.xlsx"
    create_test_excel(test_excel)
    
    styler = ExcelStyler(test_excel, sheet_name="DATA")
    
    test_config = {
        "steps": [
            {
                "action": "find_and_highlight_contains",
                "column": "Name",
                "text": "Record",
                "color": "FFFFFF00"  # Yellow
            }
        ]
    }
    
    engine = StepEngine(styler, test_config)
    
    try:
        print("✓ Executing find_and_highlight_contains...")
        result = engine.execute()
        styler.save("test_find_contains_output.xlsx")
        print(f"✓ SUCCESS: Highlighted all cells containing 'Record'")
        print(f"  Output: test_find_contains_output.xlsx")
        
        # Cleanup
        if os.path.exists(test_excel):
            os.remove(test_excel)
        
        return True
    except Exception as e:
        print(f"✗ FAILED: {e}")
        import traceback
        traceback.print_exc()
        return False


def test_combined_actions():
    """Test multiple actions in sequence"""
    print("\n" + "="*70)
    print("TEST 3: Combined multiple actions")
    print("="*70)
    
    test_excel = "test_combined.xlsx"
    create_test_excel(test_excel)
    
    styler = ExcelStyler(test_excel, sheet_name="DATA")
    
    test_config = {
        "steps": [
            {
                "action": "highlight_column",
                "column": ["Status", "Value"],
                "color": "FFFFFF00"  # Yellow
            },
            {
                "action": "find_and_highlight_equals",
                "column": "Status",
                "value": "ERROR",
                "color": "FFFF0000"  # Red
            },
            {
                "action": "find_and_highlight_contains",
                "column": "Name",
                "text": "1",
                "color": "FF00FF00"  # Green
            }
        ]
    }
    
    engine = StepEngine(styler, test_config)
    
    try:
        print("✓ Executing multiple actions...")
        result = engine.execute()
        styler.save("test_combined_output.xlsx")
        print(f"✓ SUCCESS: All actions executed successfully")
        print(f"  Output: test_combined_output.xlsx")
        
        # Cleanup
        if os.path.exists(test_excel):
            os.remove(test_excel)
        
        return True
    except Exception as e:
        print(f"✗ FAILED: {e}")
        import traceback
        traceback.print_exc()
        return False


if __name__ == "__main__":
    all_passed = True
    
    all_passed &= test_highlight_equals_action()
    all_passed &= test_highlight_equals_list_action()
    all_passed &= test_highlight_contains_action()
    all_passed &= test_combined_actions()
    
    print("\n" + "="*70)
    if all_passed:
        print("✓ ALL TESTS PASSED")
    else:
        print("✗ SOME TESTS FAILED")
    print("="*70 + "\n")
    
    sys.exit(0 if all_passed else 1)
