import sys
from PySide6.QtWidgets import QApplication
from pathlib import Path

sys.path.append(str(Path(__file__).resolve().parent.parent / "src"))
from pprint import pprint
# IMPORTA TU VENTANA

from openpyxl import load_workbook

def build_tests_from_excel(excel_path):
    """
    Lee el Excel y devuelve la lista `tests` con la misma estructura
    que ya usabas en tu código.
    """

    wb = load_workbook(excel_path, data_only=True)
    ws = wb.active

    # -------------------------
    # Detectar columnas
    # -------------------------
    headers = {}
    header_row = None

    for row in ws.iter_rows(min_row=1, max_row=20):
        for cell in row:
            if cell.value:
                headers[str(cell.value).strip()] = cell.column

        if "Test N°" in headers and "Analysis" in headers:
            header_row = cell.row
            break

    if header_row is None:
        raise ValueError("No se encontraron las columnas esperadas")

    # -------------------------
    # Función para evaluar resultado
    # -------------------------
    def get_result_from_analysis(analysis):
        if not analysis:
            return "FAIL"

        text = str(analysis).upper()

        # regla: si hay al menos un FAIL → FAIL
        if "FAIL" in text:
            return "FAIL"

        if "PASS" in text:
            return "PASS"

        return "FAIL"

    # -------------------------
    # Construcción de tests
    # -------------------------
    tests = []

    for row in ws.iter_rows(min_row=header_row + 1):

        test_num = row[headers["Test N°"] - 1].value
        if test_num == "N/A" or test_num == "n/a":
            continue

        # ignorar filas vacías
        if test_num is None:
            continue

        test_desc = ""
        if "Test Env" in headers:
            test_desc = row[headers["Test Env"] - 1].value or ""

        test_notes = ""
        if "Test scope" in headers:
            test_notes = row[headers["Test scope"] - 1].value or ""

        analysis = row[headers["Analysis"] - 1].value
        test_result = get_result_from_analysis(analysis)

        tests.append({
            "num": test_num,
            "desc": test_desc,
            "notes": test_notes,
            "result": test_result
        })

    return tests

def main():

    path = r"C:\Users\gonzpidr\Ariston Group\Alignement - General\01_INVERTERS\05_Ariston Inverter\03_Testing\02_SW Qualification - Integration Inverter-TDM\2026_04_29 TDM 70.35.00-Ariston 15.15.01\02_Report\Report_Integration_ATG_15.15.01 - TDM_70.35.00.xlsx"
    data = build_tests_from_excel(path)
    pprint(data)

    sys.exit(0)

    


if __name__ == "__main__":
    main()