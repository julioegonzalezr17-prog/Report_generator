from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Font, Alignment, Border, Side
import shutil


# =========================
# CONFIGURACIÓN
# =========================

PASSWORD = "prova"

PASS_FILL = PatternFill(
    start_color="C6EFCE",
    end_color="C6EFCE",
    fill_type="solid"
)

FAIL_FILL = PatternFill(
    start_color="FFC7CE",
    end_color="FFC7CE",
    fill_type="solid"
)

HEADER_FONT = Font(bold=True)
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center")
HEADER_BORDER = Border(
    bottom=Side(style="thin")
)


# =========================
# HELPERS
# =========================

def get_writable_cell(ws, row, column):
    """
    Devuelve una celda escribible.
    Si pertenece a un rango combinado, devuelve la celda top-left del merge.
    """
    cell = ws.cell(row=row, column=column)

    if isinstance(cell, MergedCell):
        for merged_range in ws.merged_cells.ranges:
            if (
                merged_range.min_row <= row <= merged_range.max_row
                and merged_range.min_col <= column <= merged_range.max_col
            ):
                return ws.cell(
                    row=merged_range.min_row,
                    column=merged_range.min_col
                )

    return cell


def find_table_start_row(ws, min_empty_rows=3):
    """
    Detecta automáticamente dónde iniciar la tabla de tests
    buscando un bloque de filas vacías consecutivas.
    """
    empty_count = 0

    for row in range(1, ws.max_row + 1):
        row_has_data = False

        for col in range(1, ws.max_column + 1):
            if ws.cell(row=row, column=col).value not in (None, ""):
                row_has_data = True
                break

        if row_has_data:
            empty_count = 0
        else:
            empty_count += 1
            if empty_count >= min_empty_rows:
                return row - min_empty_rows + 1

    return ws.max_row + 2

def find_value_column_after_label(ws, label_cell):
    """
    Devuelve la columna correcta donde debe escribirse el valor
    (a la derecha real del label, respetando merged cells).
    """
    row = label_cell.row
    col = label_cell.column

    # Si el label pertenece a un merged range
    for merged_range in ws.merged_cells.ranges:
        if (
            merged_range.min_row <= row <= merged_range.max_row
            and merged_range.min_col <= col <= merged_range.max_col
        ):
            return merged_range.max_col + 1

    # Label normal (no merge)
    return col + 1


# =========================
# FUNCIÓN PRINCIPAL
# =========================

def update_excel_template(
    template_path,
    output_path,
    data_path,
    header_data, 
    tests
):
    template_path = Path(template_path)
    if output_path is None:
        output_path = Path(template_path)
    else:
        output_path = Path(output_path)

    if not template_path.exists():
        raise FileNotFoundError(f"Template not found: {template_path}")

    if template_path.suffix.lower() != ".xlsx":
        raise ValueError("The template must be an .xlsx file")   
    excel_dir = template_path.parent
    print(f"Excel directory: {excel_dir}")
    print(f"Data path: {data_path}")
    shutil.copytree(data_path, excel_dir, dirs_exist_ok=True)

    report_folder = excel_dir / "02_Report"
    data_folder = excel_dir / "03_Data"


    wb = load_workbook(template_path)
    ws = wb.active

    # Desproteger hoja
    ws.protection.disable()

    # =========================
    # Rellenar campos header (merge-safe)
    # =========================
    for row in ws.iter_rows():
        for cell in row:
            if cell.value in header_data:                
                value_col = find_value_column_after_label(ws, cell)
                writable_cell = get_writable_cell(
                    ws,
                    cell.row,
                    value_col
                )
                writable_cell.value = header_data[cell.value]

    # =========================
    # Detectar inicio de la tabla
    # =========================
    start_row = find_table_start_row(ws)
    # =========================
    #  FILA "TEST REPORT FILE"
    # =========================

    label_row = start_row - 2  # deja una fila de aire

    label_cell = get_writable_cell(ws, label_row, 2)
    label_cell.value = "Test report file"
    label_cell.font = Font(bold=True)

    value_cell = get_writable_cell(ws, label_row, 3)
    value_cell.value = "Report file" if report_folder.exists() else "Report folder not found"
    value_cell.style = "Hyperlink"
    value_cell.hyperlink = str(report_folder) if report_folder.exists() else None


    start_row = find_table_start_row(ws)
    # =========================
    # Cabecera de la tabla
    # =========================
    headers = [
        "Test #",
        "Description",
        "Logs",
        "Notes",
        "Result"
    ]
    for col, header in enumerate(headers, start=2):
        cell = get_writable_cell(ws, start_row, col)
        cell.value = header
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGNMENT
        cell.border = HEADER_BORDER
    # =========================
    # Insertar tests
    # =========================
    for idx, test in enumerate(tests, start=start_row + 1):
        test_data_path = data_folder / f"Test_{test['num']}"
        if not test_data_path.exists():
            test_data_path = "No data"
        get_writable_cell(ws, idx, 2).value = test["num"]
        get_writable_cell(ws, idx, 3).value = test["desc"]
        get_writable_cell(ws, idx, 4).value = f"Log test {test['num']}" if test_data_path != "No data" else "No data"
        get_writable_cell(ws, idx, 4).style = "Hyperlink"
        get_writable_cell(ws, idx, 4).hyperlink = str(test_data_path) if test_data_path != "No data" else None
        get_writable_cell(ws, idx, 5).value = test["notes"]
        result_cell = get_writable_cell(ws, idx, 6)
        result_cell.value = test["result"]
        result_cell.fill = (
            PASS_FILL if test["result"] == "PASS" else FAIL_FILL
        )

    # Reproteger hoja
    ws.protection.set_password(PASSWORD)
    ws.protection.enable()   
    # =========================
    # GUARDADO
    # =========================

    wb.save(output_path)

