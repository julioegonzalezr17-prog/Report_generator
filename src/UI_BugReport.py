import pandas as pd
import copy
from PySide6.QtWidgets import (
    QDialog, QLabel, QPushButton, QVBoxLayout, QHBoxLayout,
    QScrollArea, QWidget, QCheckBox, QComboBox,
    QLineEdit, QTextEdit, QFileDialog, QMessageBox
)
from PySide6.QtGui import QFont, QIcon
from PySide6.QtCore import Qt
from openpyxl import load_workbook
from datetime import datetime
from openpyxl.styles import Border, Side, PatternFill
from dictDataIntegration import sw_version
import logging
from pathlib import Path

logger = logging.getLogger(__name__)

class BugReportWindow(QDialog):
    """
    Second window:
    - Receives the full data_dict from BugSelectionWindow
    - Allows selecting the buglist Excel file
    - Reports new bugs and links them to existing ones
    """

    def __init__(self, data_dict: dict):
        super().__init__()

        self.data = copy.deepcopy(data_dict)
        self.report_data = self.data.get("report_data", {})
        tests_dict = self.data.get("Test", {})       
        self.tests_with_bug = {
            k: v for k, v in tests_dict.items()
            if isinstance(v, dict) and v.get("BUG") is True
        }


        self.buglist_path = ""
        self.existing_bugs = []
        self.related_checkboxes = {}

        self.setModal(True)
        self.setWindowTitle("New Bug Report" + "Version " + sw_version)
        self.setWindowIcon(QIcon(":/monitoring.png"))
        self.resize(900, 700)

        self.setStyleSheet("""
            QWidget {
                font-size: 14px;
            }
            QPushButton {
                min-height: 40px;
                padding: 8px 16px;
                background-color: #2d89ef;
                color: white;
                font-weight: bold;
                border-radius: 6px;
            }
            QPushButton:hover {
                background-color: #1b5fbd;
            }
        """)


        # =================================================
        # Bug inputs
        # =================================================
        self.title_input = QLineEdit()
        self.title_input.setPlaceholderText("Short bug name")

        self.priority_input = QComboBox()
        self.priority_input.addItems(["Low", "Medium", "High"])

        self.description_input = QTextEdit()
        self.description_input.setPlaceholderText("Bug description / Notes Technology")

        # =================================================
        # Buglist path selector
        # =================================================
        self.path_edit = QLineEdit()
        self.path_edit.setReadOnly(True)

        btn_browse = QPushButton("Browse")
        btn_browse.clicked.connect(self._select_buglist_file)

        path_layout = QHBoxLayout()
        path_layout.addWidget(self.path_edit)
        path_layout.addWidget(btn_browse)

        # =================================================
        # Existing bugs area
        # =================================================
        self.bugs_scroll = QScrollArea()
        self.bugs_scroll.setWidgetResizable(True)

        # =================================================
        # Buttons
        # =================================================
        self.btn_cancel = QPushButton("Cancel")
        self.btn_update = QPushButton("Update Buglist")

        buttons = QHBoxLayout()
        buttons.addStretch()
        buttons.addWidget(self.btn_cancel)
        buttons.addWidget(self.btn_update)

        # =================================================
        # Main layout
        # =================================================
        main = QVBoxLayout(self)

        main.addWidget(QLabel("<b>Bug Name*</b>"))
        main.addWidget(self.title_input)

        main.addWidget(QLabel("<b>Bug Priority*</b>"))
        main.addWidget(self.priority_input)

        main.addWidget(QLabel("<b>Bug Description (Notes Technology)*</b>"))
        main.addWidget(self.description_input)

        main.addWidget(QLabel("<b>Select Buglist Excel file</b>"))
        main.addLayout(path_layout)

        main.addWidget(QLabel("<b>Existing bugs (same Inverter Model)</b>"))
        main.addWidget(self.bugs_scroll, stretch=1)

        main.addLayout(buttons)

        self.btn_cancel.clicked.connect(self.reject)
        self.btn_update.clicked.connect(self._on_update)

    # =================================================
    # Select Excel file
    # =================================================
    def _select_buglist_file(self):
        path, _ = QFileDialog.getOpenFileName(
            self,
            "Select Buglist Excel file",
            "",
            "Excel files (*.xlsx)"
        )

        if not path:
            return

        self.buglist_path = path
        self.path_edit.setText(path)
        self._load_existing_bugs()

    # =================================================
    # Load bugs from Excel
    # =================================================
    def _load_existing_bugs(self):
        self.existing_bugs.clear()
        self.related_checkboxes.clear()

        INVERTER_MODEL = self.report_data.get("inverter_model")
        xls = pd.ExcelFile(self.buglist_path)

        for sheet in xls.sheet_names:
            df = xls.parse(sheet)
            
            for _, row in df.iterrows():
                self.existing_bugs.append({
                    "item": row.get("Item n", ""),
                    "issue": row.get("ISSUES", ""),
                    "sheet": sheet
                })
        # Show only bugs for current machine
        visible_bugs = [
            b for b in self.existing_bugs
            if b["sheet"] == INVERTER_MODEL
        ]

        self._update_bug_list_ui(visible_bugs)

    # =================================================
    # Build existing bug list UI
    # =================================================
    def _update_bug_list_ui(self, bugs):
        
        container = QWidget()

        font = container.font()
        if font.pointSize() <= 0:
            font.setPointSize(10)
        container.setFont(font)

        layout = QVBoxLayout(container)
        if not bugs:
            lbl = QLabel("No existing bugs found for this machine ✅")
            lbl.setStyleSheet("color: green;")
            layout.addWidget(lbl)
        else:
            for bug in bugs:
                text = f"{bug['item']} - {bug['issue']}"
                cb = QCheckBox(text)
                self.related_checkboxes[bug["item"]] = cb
                layout.addWidget(cb)

        layout.addStretch()
        self.bugs_scroll.setWidget(container)

    # =================================================
    # Update Excel
    # =================================================
    def _on_update(self):
        if not self.buglist_path:
            QMessageBox.warning(
                self,
                "No bugs list file",
                "Please select a bug list Excel file to proceed."
            )
            return
        if not self.title_input.text().strip():
            QMessageBox.warning(
                self,
                "Missing bug name",
                "Please enter a short name for the bug to proceed."
            )
            return
        if not self.description_input.toPlainText().strip():
            QMessageBox.warning(
                self,
                "Missing bug description",
                "Please enter a description for the bug to proceed."
            )
            return

        self._append_bug_to_excel()
        self.accept()

    
    def _append_bug_to_excel(self):

        report_data = self.data.get("report_data", {})
        tests_dict = self.data.get("Test", {})

        INVERTER_MODEL = report_data.get("inverter_model")

        wb = load_workbook(self.buglist_path)
        ws = wb[INVERTER_MODEL]

        # =========================
        # 1️⃣ HEADERS
        # =========================
        headers = {}
        for col in range(1, ws.max_column + 1):
            val = ws.cell(row=1, column=col).value
            if val:
                headers[str(val).strip()] = col

        # =========================
        # 🔹 BUILD MULTI TEST DATA
        # =========================
        
        selected_tests = [
            (k, v) for k, v in tests_dict.items()
            if isinstance(v, dict) and v.get("BUG")
        ]

        if not selected_tests:
            print("⚠️ No BUG selected")
            return

        test_scope_text = []
        test_result_text = []

        for test_id, test_data in selected_tests:

            scope = test_data.get("test_scope", "")
            analysis = test_data.get("data_print_report", {}).get("Analysis", "")

            test_scope_text.append(f"Test {test_id}: {scope}")
            test_result_text.append(f"Test {test_id}:\n{analysis}")

        # =========================
        # 3️⃣ EVITAR DUPLICADO
        # =========================
        # Revisar si ya existe el mismo ISSUE (muy útil)
        issue_text = self.title_input.text()

        for r in range(2, ws.max_row + 1):
            cell_val = ws.cell(row=r, column=headers.get("ISSUES")).value
            if cell_val == issue_text:
                print("⚠️ Bug already exists → skip")
                return

        # =========================
        # 4️⃣ NUEVA FILA
        # =========================
        new_row_idx = ws.max_row + 1

        # =========================
        # 5️⃣ RELATED BUGS
        # =========================
        related = [
            str(item)
            for item, cb in self.related_checkboxes.items()
            if cb.isChecked()
        ]

        # =========================
        # 6️⃣ ITEM NUMBER
        # =========================
        items = [
            ws.cell(row=r, column=headers["Item n"]).value
            for r in range(2, ws.max_row + 1)
            if ws.cell(row=r, column=headers["Item n"]).value
        ]

        item_number = max(items) + 1 if items else 1

        # =========================
        # 7️⃣ WRITE HELPER
        # =========================
        def write(col_name, value):
            if col_name in headers:
                ws.cell(row=new_row_idx, column=headers[col_name]).value = value

        # =========================
        # 8️⃣ WRITE DATA
        # =========================
        write("Item n", item_number)
        write("ISSUES", issue_text)

        write("Test scope", "\n".join(test_scope_text))
        write("Test result", "\n\n".join(test_result_text))
        write("Priority", self.priority_input.currentText())
        
        write("Notes Technology", self.description_input.toPlainText())

        # ✅ STATUS
        write("Status", "OPEN")

        write("Related BUGs", ", ".join(related))

        write("DSP version", ".".join(report_data.get("inverter_dsp", [])))
        write("EEPROM version", ".".join(report_data.get("inverter_eeprom", [])))
        write("ARM version", ".".join(report_data.get("inverter_arm", [])))
        write("CB version", ".".join(report_data.get("control_board", [])))
        write("TDM version", ".".join(report_data.get("tdm_version", [])))

        write("Date Tested", datetime.now().strftime("%d/%m/%Y"))

        # =========================
        # 9️⃣ MULTI LOGS ✅
        # =========================
        if "Logs" in headers:

            base_path = Path(self.data.get("RDP_path", ""))
            logs_cell = ws.cell(row=new_row_idx, column=headers["Logs"])

            logs_texts = []
            first_link_set = False

            for test_id, _ in selected_tests:

                log_path = base_path / "03_Data" / f"Test_{test_id}"

                text = f"Logs_test_{test_id}"
                logs_texts.append(text)

                # Excel solo acepta 1 hyperlink real → usamos el primero
                if log_path.exists() and not first_link_set:
                    logs_cell.hyperlink = str(log_path)
                    first_link_set = True

            logs_cell.value = "\n".join(logs_texts)
            logs_cell.style = "Hyperlink"

        # =========================
        # 🔟 BORDES + COLOR STATUS ✅
        # =========================
        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin")
        )

        RED_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        GREEN_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
        YELLOW_FILL = PatternFill(start_color="FFFFEB9C",end_color="FFFFEB9C",fill_type="solid")        
        ORANGE_FILL = PatternFill(start_color="FFF4B084",end_color="FFF4B084",fill_type="solid")


        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row=new_row_idx, column=col)
            cell.border = thin_border

            # ✅ COLOR STATUS
            if col == headers.get("Priority"):
                if self.priority_input.currentText() == "High":
                    cell.fill = RED_FILL
                elif self.priority_input.currentText() == "Medium":
                    cell.fill = ORANGE_FILL
                elif self.priority_input.currentText() == "Low":
                    cell.fill = YELLOW_FILL
            if col == headers.get("Status"):
                if cell.value == "OPEN":
                    cell.fill = RED_FILL
                elif cell.value == "CLOSED":   # por si en futuro
                    cell.fill = GREEN_FILL

        # =========================
        # SAVE
        # =========================
        ws.freeze_panes = "A2"
        wb.save(self.buglist_path)

    
    
